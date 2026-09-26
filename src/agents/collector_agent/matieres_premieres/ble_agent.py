from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from src.schemas import MatierePremiereOutput, MarketDataPoint


def run() -> MatierePremiereOutput:
    """Charge le benchmark international mensuel Wheat, US HRW du Pink Sheet."""
    workbook_path = (
        Path(__file__).resolve().parents[4]
        / "data"
        / "raw"
        / "worldbank_commodity_prices_monthly.xlsx"
    )
    if not workbook_path.exists():
        return MatierePremiereOutput(
            matiere="ble",
            points=[],
            source=f"Fichier introuvable: {workbook_path}",
            is_synthetic=True,
        )

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Monthly Prices"]
    header_row = None
    columns = {}
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        normalized = [" ".join(str(value).split()) if value is not None else "" for value in row]
        if "Wheat, US HRW" in normalized:
            header_row = row_number
            columns = {name: normalized.index(name) for name in normalized if name}
            break

    if header_row is None:
        workbook.close()
        raise ValueError("Colonne 'Wheat, US HRW' absente du classeur World Bank.")

    price_column = columns["Wheat, US HRW"]
    brent_column = columns.get("Crude oil, Brent")
    urea_column = columns.get("Urea")
    points = []
    for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
        period = row[0]
        price = row[price_column]
        if not isinstance(period, str) or not isinstance(price, (int, float)):
            continue
        try:
            year, month = period.split("M")
            observation_date = date(int(year), int(month), 1)
        except (ValueError, TypeError):
            continue
        if price <= 0:
            continue
        brent = row[brent_column] if brent_column is not None else None
        urea = row[urea_column] if urea_column is not None else None
        points.append(
            MarketDataPoint(
                date=observation_date,
                code_sh="1001",
                prix_reference_usd_tonne=float(price),
                brent_usd_baril=float(brent) if isinstance(brent, (int, float)) and brent > 0 else None,
                uree_usd_tonne=float(urea) if isinstance(urea, (int, float)) and urea > 0 else None,
            )
        )

    workbook.close()
    return MatierePremiereOutput(
        matiere="ble",
        points=points,
        source="World Bank Pink Sheet - Wheat, US HRW (USD/tonne)",
        is_synthetic=False,
    )
